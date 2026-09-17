from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path('D:/context/忆安智护_个人工作答辩版.pptx')
CHARTS = Path('D:/context/context-keeper-main/experiments/results/reports/charts')

NAVY = RGBColor(6, 29, 78)
PANEL = RGBColor(9, 54, 116)
PANEL2 = RGBColor(10, 57, 121)
CYAN = RGBColor(50, 217, 242)
TEAL = RGBColor(36, 199, 180)
AMBER = RGBColor(244, 198, 78)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(185, 214, 238)
LINE = RGBColor(59, 169, 219)


def set_run(run, size, color=WHITE, bold=False, font='Microsoft YaHei'):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_text(slide, text, x, y, w, h, size=14, color=WHITE, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, font='Microsoft YaHei'):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    set_run(run, size, color, bold, font)
    return shape


def rect(slide, x, y, w, h, fill, line=None, radius=False, transparency=0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.color.rgb = line or fill
    shape.line.width = Pt(1)
    return shape


def add_background(slide, section, index):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    rect(slide, 0, 0, 13.333, 0.11, CYAN)
    rect(slide, 0, 6.99, 13.333, 0.51, RGBColor(4, 21, 56))
    add_text(slide, section.upper(), 0.58, 0.31, 3.3, 0.22, 8.5, CYAN, True, font='Arial')
    add_text(slide, '忆安智护 | 个人负责工作答辩', 0.58, 7.13, 4.2, 0.16, 8, MUTED)
    add_text(slide, f'{index:02d}', 12.05, 7.08, 0.68, 0.22, 12, CYAN, True, PP_ALIGN.RIGHT, font='Arial')


def add_title(slide, text, subtitle):
    add_text(slide, text, 0.58, 0.72, 11.9, 0.55, 27, WHITE, True)
    if subtitle:
        add_text(slide, subtitle, 0.6, 1.31, 11.7, 0.26, 11.5, MUTED)


def add_card(slide, x, y, w, h, heading, body, accent, number=None):
    rect(slide, x, y, w, h, PANEL, accent, True, 7)
    if number:
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.27), Inches(y + 0.3), Inches(0.6), Inches(0.6))
        circle.fill.solid(); circle.fill.fore_color.rgb = accent; circle.line.color.rgb = accent
        add_text(slide, str(number), x + 0.27, y + 0.42, 0.6, 0.18, 13, NAVY, True, PP_ALIGN.CENTER, font='Arial')
        title_x = x + 1.02
    else:
        title_x = x + 0.28
    add_text(slide, heading, title_x, y + 0.3, w - (title_x - x) - 0.25, 0.28, 16, WHITE, True)
    add_text(slide, body, x + 0.34, y + 1.12, w - 0.66, h - 1.35, 11.5, MUTED, False, PP_ALIGN.LEFT, MSO_ANCHOR.TOP)


def add_metric(slide, x, y, label, value, note, accent):
    rect(slide, x, y, 2.75, 1.52, PANEL2, accent, True, 4)
    add_text(slide, label, x + 0.14, y + 0.17, 2.47, 0.23, 11.5, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, value, x + 0.12, y + 0.55, 2.51, 0.35, 24, accent, True, PP_ALIGN.CENTER, font='Arial')
    add_text(slide, note, x + 0.12, y + 1.13, 2.51, 0.18, 8.5, MUTED, False, PP_ALIGN.CENTER)


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# 1. Cover
s = prs.slides.add_slide(blank)
s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
rect(s, 0, 0, 13.333, 0.12, CYAN)
ring = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.75), Inches(1.0), Inches(2.0), Inches(2.0))
ring.fill.background(); ring.line.color.rgb = CYAN; ring.line.width = Pt(2)
add_text(s, '忆安智护', 0.72, 1.38, 5.8, 0.67, 35, WHITE, True)
add_text(s, '个人负责工作答辩', 0.74, 2.2, 4.7, 0.35, 19, CYAN, True)
add_text(s, '代码编写与修改 · 实验运行与审核 · PPT与策划书制作', 0.75, 3.0, 7.4, 0.32, 15, MUTED)
rect(s, 0.75, 3.67, 5.58, 0.025, CYAN)
add_text(s, '郑州大学网络空间安全学院\n答辩人：宫正行', 0.75, 4.22, 4.8, 0.75, 14, WHITE)
add_text(s, 'CODE\nEXPERIMENT\nMATERIAL', 9.3, 3.42, 2.9, 1.18, 15, AMBER, True, PP_ALIGN.CENTER, font='Arial')
add_text(s, '01', 0.75, 6.7, 0.6, 0.22, 13, CYAN, True, font='Arial')

# 2. Scope
s = prs.slides.add_slide(blank)
add_background(s, 'PERSONAL SCOPE', 2)
add_title(s, '我主要负责的三项工作', '以工程实现、可信实验和答辩材料为主线推进项目落地')
add_card(s, 0.62, 1.92, 3.85, 3.82, '代码编写与修复', '围绕认证、检索、因果、安全与遗忘链路进行定位和修改。\n\n重点处理：JWT 接入、配置切换、三路 RRF、跨存储删除与错误分母。', CYAN, 1)
add_card(s, 4.74, 1.92, 3.85, 3.82, '实验运行与审核', '搭建固定数据集、评测脚本、运行清单和原始结果审核流程。\n\n确保 API 错误不进入业务指标，结果可回溯至数据、配置与日志。', TEAL, 2)
add_card(s, 8.86, 1.92, 3.85, 3.82, 'PPT与策划书制作', '将技术逻辑、场景价值和实测结果转化为答辩材料。\n\n制作项目叙事、实验页、图表、报告与现场展示材料。', AMBER, 3)

# 3. Engineering
s = prs.slides.add_slide(blank)
add_background(s, 'ENGINEERING', 3)
add_title(s, '代码工作：把“能运行”变成“可验证”', '我负责的修改重点，是让接口、数据范围和结果口径真正对齐')
engineering = [
    ('认证与路由', '统一 JWT 登录和受保护 API，避免 404、认证失败被误算为业务成功。', CYAN),
    ('检索链路', '补齐 doc_id 溯源、会话隔离和直接三路 RRF 评测路径。', TEAL),
    ('安全与因果', '补充输入策略，修复 PCCM 置信度融合与中文 O-M-P-R 解析。', AMBER),
    ('遗忘验证', '实现 Qdrant、TimescaleDB、Neo4j、缓存的跨存储计数与删除校验。', CYAN),
]
for i, (heading, body, accent) in enumerate(engineering):
    x = 0.78 if i % 2 == 0 else 6.87
    y = 1.9 if i < 2 else 4.08
    add_card(s, x, y, 5.63, 1.63, heading, body, accent, i + 1)

# 4. Experiment
s = prs.slides.add_slide(blank)
add_background(s, 'EXPERIMENT', 4)
add_title(s, '实验工作：固定数据、独立配置、原始证据', '把一次性展示转化为可复跑、可检查、可解释的实验过程')
rows = [
    ['实验', '固定数据', '我负责的验证方式'],
    ['检索', '30条语料 + 50条标注查询', '3配置分别运行，保存 MRR、P@5、R@5 与来源 trace'],
    ['因果', '20条人工 O-M-P-R 标注', '保存有效响应、严格四元组匹配、置信度与延迟'],
    ['安全', '6类攻击，每类20条', 'API 错误单列；ASR 仅按有效响应计算'],
    ['遗忘', '目标会话 + 对照会话', '3次独立试验，跨四类存储计数与检索探针验证'],
]
table = s.shapes.add_table(len(rows), 3, Inches(0.75), Inches(1.85), Inches(11.82), Inches(3.55)).table
table.columns[0].width = Inches(1.25); table.columns[1].width = Inches(3.15); table.columns[2].width = Inches(7.42)
for r, values in enumerate(rows):
    for c, value in enumerate(values):
        cell = table.cell(r, c)
        cell.text = value
        cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor(20, 112, 194) if r == 0 else PANEL
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for paragraph in cell.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.CENTER
            for run in paragraph.runs:
                set_run(run, 11 if r else 12, WHITE, r == 0)
rect(s, 0.75, 5.73, 11.82, 0.62, RGBColor(8, 51, 107), CYAN, True, 4)
add_text(s, '关键原则：原始 JSON + 配置哈希 + 数据集哈希 + 运行日志；认证、超时和服务错误不进入业务指标分母。', 1.0, 5.92, 11.3, 0.2, 11, AMBER, True, PP_ALIGN.CENTER)

# 5. Results
s = prs.slides.add_slide(blank)
add_background(s, 'MEASURED RESULTS', 5)
add_title(s, '实测结果：我的工作如何形成可答辩证据', '以下均来自通过有效性门槛的原始 JSON 与运行清单')
add_metric(s, 0.78, 1.78, '三路检索 MRR', '0.491', '50条查询 | 103.5 ms', CYAN)
add_metric(s, 3.78, 1.78, '因果严格 O-M-P-R', '85%', '20条标注案例', TEAL)
add_metric(s, 6.78, 1.78, '安全防御成功率', '93.3%', '120条攻击样本', AMBER)
add_metric(s, 9.78, 1.78, '机器遗忘', '3 / 3', '跨存储试验通过', CYAN)
add_card(s, 0.78, 3.72, 5.74, 2.22, '检索优化结果', '通过跳过评测场景下的 LLM 语义分析与回答生成，将三路 RRF 检索从约 16.4 秒降至 103.5 毫秒，同时保留 vector、timeline、knowledge 的来源审计。', CYAN)
chart = CHARTS / 'retrieval_quality.png'
if chart.exists():
    s.shapes.add_picture(str(chart), Inches(6.85), Inches(3.72), Inches(5.7), Inches(2.22))

# 6. Materials and close
s = prs.slides.add_slide(blank)
add_background(s, 'MATERIAL & CLOSING', 6)
add_title(s, '材料制作与答辩交付', '把工程过程转换为评委能够快速理解和复核的表达')
deliverables = [
    ('PPT叙事', '将项目背景、技术架构、实验结果和未来方向串成答辩主线。'),
    ('策划书', '梳理需求、系统设计、实现、测试与实习工作内容。'),
    ('实验报告', '生成 HTML 报告、CSV、PNG 图表和原始 JSON 证据。'),
]
for i, (heading, body) in enumerate(deliverables):
    y = 1.82 + i * 1.24
    rect(s, 0.82, y, 7.06, 0.94, PANEL, AMBER if i == 2 else CYAN, True, 5)
    add_text(s, f'0{i + 1}', 1.08, y + 0.31, 0.44, 0.18, 13, AMBER if i == 2 else CYAN, True, font='Arial')
    add_text(s, heading, 1.82, y + 0.22, 1.35, 0.22, 14, WHITE, True)
    add_text(s, body, 3.2, y + 0.22, 4.25, 0.4, 10.5, MUTED)
add_card(s, 8.28, 1.82, 4.22, 3.42, '个人工作总结', '我承担的核心不是单纯展示功能，而是把代码修复、实验验证和答辩材料连接起来。\n\n让系统可运行，让指标可解释，让结果可复核。', TEAL)
add_text(s, '谢谢各位老师', 8.77, 4.55, 3.1, 0.25, 15, AMBER, True, PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)

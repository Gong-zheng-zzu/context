from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path("D:/context")
SOURCE = next(path for path in ROOT.glob("*.pptx") if "v2" not in path.stem)
OUTPUT = ROOT / "忆安智护实验修改版_v2.pptx"

WHITE = RGBColor(255, 255, 255)
CYAN = RGBColor(54, 226, 255)
NAVY = RGBColor(8, 42, 104)
PANEL = RGBColor(9, 61, 130)
TEAL = RGBColor(41, 205, 202)
GOLD = RGBColor(255, 204, 73)


def set_text(shape, text, size=14, color=WHITE, bold=False, align=PP_ALIGN.CENTER):
    frame = shape.text_frame
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return frame


def set_table(table, values, header_rows=1, size=10.5):
    for row_index, row_values in enumerate(values):
        for col_index, value in enumerate(row_values):
            if col_index >= len(table.columns) or row_index >= len(table.rows):
                continue
            cell = table.cell(row_index, col_index)
            cell.text = str(value)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"
                    run.font.size = Pt(size + (0.5 if row_index < header_rows else 0))
                    run.font.bold = row_index < header_rows
                    run.font.color.rgb = WHITE


def remove_shape(shape):
    element = shape._element
    element.getparent().remove(element)


def add_metric_card(slide, left, top, width, title, value, note, accent):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(1.46))
    card.fill.solid()
    card.fill.fore_color.rgb = PANEL
    card.line.color.rgb = accent
    card.line.width = Pt(1.2)
    frame = card.text_frame
    frame.clear()
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    frame.margin_top = Inches(0.08)
    p1 = frame.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = title
    r1.font.name = "Microsoft YaHei"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    p2 = frame.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = value
    r2.font.name = "Arial"
    r2.font.size = Pt(23)
    r2.font.bold = True
    r2.font.color.rgb = accent
    p3 = frame.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    r3 = p3.add_run()
    r3.text = note
    r3.font.name = "Microsoft YaHei"
    r3.font.size = Pt(8.5)
    r3.font.color.rgb = WHITE


def update_slide_18(slide):
    set_text(slide.shapes[1], "实验设计与验证范围", size=25, bold=True, align=PP_ALIGN.LEFT)
    set_text(slide.shapes[7], "五步验证流程", size=16, bold=True, align=PP_ALIGN.LEFT)
    flow_texts = ["固定\n数据集", "独立\n配置", "JWT/API\n预检", "原始\nJSON", "图表与\n报告"]
    for shape_index, value in zip(range(31, 36), flow_texts):
        set_text(slide.shapes[shape_index], value, size=11, bold=True)

    set_table(
        slide.shapes[5].table,
        [
            ["实验", "固定数据", "正式评测", "结果"],
            ["检索", "30条护理语料\n50条标注查询", "3配置 x 50查询", "完成"],
            ["因果", "20条 O-M-P-R 标注", "20个案例", "完成"],
            ["安全", "6类攻击样本", "120条", "完成"],
            ["遗忘", "30条跨存储数据\n对照会话", "3次独立试验", "完成"],
        ],
        size=9.4,
    )
    set_text(slide.shapes[20], "四类实验与固定数据集", size=16, color=CYAN, bold=True, align=PP_ALIGN.LEFT)
    set_text(slide.shapes[19], "三类检索查询：时间 17｜因果 17｜通用 16", size=11, color=CYAN, bold=True, align=PP_ALIGN.LEFT)


def update_slide_19(slide):
    config_table = slide.shapes[4]
    old_chart = slide.shapes[7]
    chart_title = slide.shapes[8]
    causal_table = slide.shapes[10]
    causal_title = slide.shapes[11]
    set_text(slide.shapes[1], "三路检索与因果抽取实验结果", size=25, bold=True, align=PP_ALIGN.LEFT)
    set_text(slide.shapes[2], "RETRIEVAL & CAUSALITY", size=10, color=CYAN, bold=True, align=PP_ALIGN.LEFT)
    set_table(
        config_table.table,
        [
            ["配置", "向量检索", "可信过滤", "时序检索", "图谱检索"],
            ["Naive RAG", "是", "否", "否", "否"],
            ["RAG + Filter", "是", "是", "否", "否"],
            ["Full System", "是", "是", "是", "是"],
        ],
        size=11,
    )

    set_text(chart_title, "50条检索实验", size=15, color=CYAN, bold=True, align=PP_ALIGN.LEFT)
    remove_shape(old_chart)
    add_metric_card(slide, 6.72, 2.04, 1.88, "MRR", "0.491", "Naive / Filter: 0.405", TEAL)
    add_metric_card(slide, 8.83, 2.04, 1.88, "Precision@5", "0.164", "Full System", GOLD)
    add_metric_card(slide, 10.94, 2.04, 1.88, "Recall@5", "0.582", "Avg latency: 103.5 ms", CYAN)

    set_text(causal_title, "因果结果（20例）", size=13, color=CYAN, bold=True, align=PP_ALIGN.LEFT)
    set_table(
        causal_table.table,
        [
            ["指标", "实测值", "说明"],
            ["关系发现率", "95%", "19 / 20"],
            ["严格 O-M-P-R", "85%", "17 / 20"],
        ],
        size=10,
    )


def update_slide_20(slide):
    set_text(slide.shapes[2], "安全防护与跨存储遗忘验证", size=25, bold=True, align=PP_ALIGN.LEFT)
    set_text(slide.shapes[1], "安全结果按有效 API 响应统计；遗忘结果按三次独立跨存储试验验证。", size=11, color=GOLD, bold=True, align=PP_ALIGN.LEFT)
    set_table(
        slide.shapes[15].table,
        [
            ["安全验证", "实测值", "结论"],
            ["API有效响应", "120 / 120", "指标分母完整"],
            ["攻击成功率", "6.7%", "存在8条漏防"],
            ["防御成功率", "93.3%", "6类攻击评测"],
            ["输入链一致性", "12 / 12", "轻量与聊天预生成"],
        ],
        size=9.8,
    )
    set_text(slide.shapes[3], "三次遗忘：跨存储计数验证", size=14, color=CYAN, bold=True, align=PP_ALIGN.LEFT)
    set_table(
        slide.shapes[18].table,
        [
            ["目标存储", "遗忘前 → 遗忘后"],
            ["Qdrant", "30 → 0"],
            ["TimescaleDB", "30 → 0"],
            ["Neo4j", "349 → 0"],
            ["缓存/会话", "3 → 0"],
        ],
        size=10,
    )
    set_text(slide.shapes[17], "对照会话计数保持不变｜检索探针无目标文档残留", size=11, color=GOLD, bold=True)


prs = Presentation(SOURCE)
update_slide_18(prs.slides[17])
update_slide_19(prs.slides[18])
update_slide_20(prs.slides[19])
prs.save(OUTPUT)
print(OUTPUT)

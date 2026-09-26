from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT.parent / "deliverables" / "策划书-强网杯实验增强版-20260926.docx"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_run_font(r)
        p.paragraph_format.space_after = Pt(3)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = cell.paragraphs[0].add_run(header)
        set_run_font(r, size=9, bold=True, color="FFFFFF")
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_border(cell)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, "F2F6FA")
            r = cell.paragraphs[0].add_run(str(value))
            set_run_font(r, size=9)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            set_cell_border(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size in (("Title", 23), ("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("忆安智护 强网杯实验与策划书增强版")
    set_run_font(r, size=23, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("面向智慧养老的护理数据安全治理与可信辅助决策平台")
    set_run_font(r, size=12, color="1F4E78")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("版本 2026-09-26  |  实验方案与证据门禁")
    set_run_font(r, size=9, color="666666")

    doc.add_heading("一 项目定位与比赛主线", level=1)
    add_para(doc, "本项目面向智慧养老数字化升级中的护理数据治理难题：护理记录碎片化、健康信息敏感、角色权限复杂、辅助结论难追溯以及用户数据难以撤回。项目不替代护工或医生，不输出诊断和处方，而是把护理记录变成可审计、可解释、可检索、可删除的受控记忆。")
    add_para(doc, "比赛现场只展示一条可复现主链：合成或合法脱敏护理记录 -> ASDF/CASIA/PCCM-S 安全判定 -> O-M-P-R 因果证据 -> 人工纠错 -> 三存储写入 -> RRF 审计检索 -> 用户级删除证明。")
    add_para(doc, "大模型仅作为受约束的候选裁决器。联网研究 Agent 只负责白名单资料的结构化整理，必须人工确认，不参与安全、因果和 RRF 正式指标。")

    doc.add_heading("二 核心技术与可信边界", level=1)
    add_table(doc, ["技术模块", "比赛中展示的能力", "当前证据边界"], [
        ("ASDF", "对抗输入归一化、可逆 span 映射、重检测", "需使用固定攻击集和正常样本重新评测"),
        ("CASIA", "文本上下文窗口、正负关键词、误伤抑制", "定位为上下文误伤抑制器，不宣称多模态"),
        ("PCCM-S", "规则、词典、上下文和安全模型的证据校准", "固定配置先验收，校准参数只用开发集"),
        ("O-M-P-R", "对象、媒介、属性、结果及原文证据 span", "历史严格 Tuple F1 约 0.513，需冻结集复测"),
        ("RRF", "三路来源、排名、分数、延迟和降级审计", "30 条查询提升为受控初步结果，需扩展人工真值"),
        ("删除验证", "跨 Qdrant、Neo4j、TimescaleDB 和缓存的删除证明", "称用户/会话级删除验证，不称参数级机器遗忘"),
    ])

    doc.add_heading("三 实验设计", level=1)
    doc.add_heading("3.1 安全防御与上下文误报抑制", level=2)
    add_para(doc, "建立六类攻击样本：空格或分隔符、同音替换、中文数字、Base64、语义绕过和上下文指代；另设不少于 100 条正常护理记录与不少于 100 条上下文歧义记录。对照组为 regex_only、asdf、asdf_casia、asdf_casia_pccm。LLM 兜底只作为能力上限，不替代确定性主链。")
    add_para(doc, "指标包括分攻击类型防御率、PII span Precision/Recall/F1、正常护理误伤率、Wilson 95% 区间、McNemar 配对检验、p50/p95 延迟、命中证据、归一化步骤和 span 映射。正式门槛为总防御率不低于 95%、各攻击类型不低于 90%、正常护理误伤率不高于 3%、PII span F1 不低于 0.92。")
    doc.add_heading("3.2 O-M-P-R 因果抽取", level=2)
    add_para(doc, "建立训练 600、开发 100、冻结测试 100 的分层数据集，按模板族切分，覆盖药物副作用、跌倒、压疮、感染、否定、时间顺序、多因果和无因果记录。两名成员独立标注，冲突由第三人仲裁。")
    add_para(doc, "比较规则、Qwen2.5-3B、规则加 PCCM-C、Ollama 3B 裁决和通过冻结门禁后的 QLoRA 3B。每条结果保留字段原文 span、候选关系、模型层级、否定/时序校验、拒答原因和人工纠错记录。正式门槛为严格 Tuple F1 不低于 0.75、字段宏平均 F1 不低于 0.88、无因果误报率不高于 5%、否定和时序准确率分别不低于 95%、在线 p95 不高于 3 秒。")
    doc.add_heading("3.3 三路 RRF 检索", level=2)
    add_para(doc, "使用同一语料、同一 top-k、同一清理流程和人工冻结真值，查询集至少 300 条，优先扩展到 500 条。比较向量、向量加过滤和真实三路 RRF；来源为空、超时或异常时，结果必须如实标记二路或向量降级。")
    add_para(doc, "每条结果固定记录唯一 doc_id、来源、各路排名、RRF 分数、来源状态、单路延迟和总墙钟延迟。正式门槛为 MRR 相对向量基线提升至少 5%，Recall@5 不下降。")
    doc.add_heading("3.4 跨存储删除证明", level=2)
    add_para(doc, "以目标用户/会话和控制用户/会话为对象，记录删除前计数、跨存储删除、删除后计数、检索探针、第二次幂等删除、控制数据保留和哈希链审计结果。统一称为“隔离用户/会话跨存储删除验证”。")
    doc.add_heading("3.5 端到端比赛闭环", level=2)
    add_para(doc, "现场依次输入合成护理记录、查看安全判定、查看 O-M-P-R 证据、人工修改并提交复核、写入三存储、执行 RRF 审计检索、执行用户级删除并导出证明。现场单次结果只用于展示闭环，不替代冻结集统计成绩。")

    doc.add_heading("四 证据登记和数据治理", level=1)
    add_table(doc, ["证据等级", "判定条件", "允许写入策划书的表达"], [
        ("已验证结果", "固定数据、配置、代码版本、原始输出和分母可复核", "当前实验结果"),
        ("受控初步结果", "样本规模较小或真值尚未冻结", "可行性或初步观察"),
        ("历史问题基线", "优化前或旧路径运行记录", "问题发现和优化动机"),
    ])
    add_para(doc, "所有正式结果必须附数据集版本与 SHA-256、规则/Prompt/模型版本、配置指纹、容器镜像、随机种子、运行时间、清理流程和失败样本。仅使用公开合成或合法脱敏数据，不在仓库提交真实护理原文、令牌、密钥、模型权重或运行目录。")
    add_para(doc, "当前历史结果单列为基线：安全纯正样本的 94.22% 是检出率，不是准确率或 F1；30 条查询的 RRF MRR 相对提升 20.5% 是受控小规模结果；因果严格 Tuple F1 约 0.513、平均延迟约 10.6 秒是历史问题基线；删除实验存在删除后检索残留，只能称跨存储删除验证。")

    doc.add_heading("五 研究 Agent 的受控扩展", level=1)
    add_para(doc, "研究 Agent 只读白名单来源，返回标题、链接、来源类型、抓取时间、摘要和引用片段；不直接写入护理事实，不生成诊断或治疗建议。联网不可用时切换到离线资料库，人工确认后才能形成护理知识摘要。该能力作为创新展示，不计入核心安全、因果和 RRF 成绩。")

    doc.add_heading("六 六周实施计划", level=1)
    add_table(doc, ["周次", "交付重点", "验收证据"], [
        ("第 1 周", "固化数据、配置、trace_id、哈希和历史基线", "实验配置、数据清单、基线原始 JSON"),
        ("第 2 周", "完成安全样本标注和 ASDF/CASIA/PCCM-S 消融", "分层结果、置信区间、误伤分析"),
        ("第 3 周", "完成因果候选、证据校验、三级回退和纠错队列", "冻结测试集前的开发集报告"),
        ("第 4 周", "完成三存储、RRF 故障降级和删除证明", "来源审计、删除前后探针、哈希证明"),
        ("第 5 周", "完成单页时间线控制台和浏览器验证", "桌面/移动截图、错误状态验证"),
        ("第 6 周", "冻结评测、Docker/Linux 门禁和发布快照", "HTML/JSON/CSV 报告、安全扫描、Git 历史检查"),
    ])

    doc.add_heading("七 发布门禁", level=1)
    add_para(doc, "每个独立阶段先补测试、运行专项门禁、记录产物，再创建独立 Git commit。最终执行：")
    code = doc.add_paragraph()
    for line in [
        "go mod verify",
        "go vet -tags http ./...",
        "go test -tags http ./...",
        "python -m unittest discover -s experiments/scripts -p \"test_*.py\"",
        "python -m unittest discover -s experiments/training -p \"test_*.py\"",
        "docker compose config --quiet",
        "docker compose build context-keeper",
        "python experiments/scripts/smoke_test.py --mode live",
        "python experiments/scripts/release_guard.py",
    ]:
        r = code.add_run(line + "\n")
        set_run_font(r, name="Consolas", size=9)

    doc.add_heading("八 明确不宣称事项", level=1)
    add_bullets(doc, [
        "冻结实验完成前，不宣称 95% 以上准确率、因果无幻觉或在线 p95 小于 3 秒。",
        "删除链路不宣称梯度投影模型遗忘、参数级机器遗忘、物理粉碎或绝对不可恢复。",
        "CASIA 当前是文本上下文判定，不宣称音频、图像或人脸能力已进入比赛主链。",
        "未经过数据集、配置和审计门禁的失败结果不得隐藏；指标未达标时保留负结果。",
    ])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("忆安智护 | 强网杯实验与策划书增强版 | 2026-09-26")
    set_run_font(r, size=8, color="666666")
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()

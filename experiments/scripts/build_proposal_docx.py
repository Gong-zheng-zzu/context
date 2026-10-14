#!/usr/bin/env python3
"""Render the competition Markdown proposal into a versioned DOCX."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size, color in (("Title", 24, "0F2F3D"), ("Heading 1", 16, "0F2F3D"), ("Heading 2", 13, "128C8C"), ("Heading 3", 11, "1B3745")):
        style = document.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    if "Code" not in [style.name for style in document.styles]:
        code = document.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        code.font.size = Pt(9)
        code.font.color.rgb = RGBColor(70, 80, 85)


def add_bullet(document: Document, text: str, numbered: bool = False) -> None:
    paragraph = document.add_paragraph(style="List Number" if numbered else "List Bullet")
    paragraph.add_run(text)


def render(markdown: Path, output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    configure_styles(document)
    document.core_properties.title = "忆安智护：强网杯报名版项目策划书"
    document.core_properties.subject = "智慧养老护理数据安全治理与可信辅助决策"
    document.core_properties.author = "忆安智护项目组"

    lines = markdown.read_text(encoding="utf-8").splitlines()
    i = 0
    first_title = True
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            paragraph = document.add_paragraph(style="Code")
            paragraph.paragraph_format.left_indent = Cm(0.5)
            paragraph.add_run("\n".join(code_lines))
            i += 1
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title" if first_title else "Heading 1")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if first_title else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.add_run(line[2:].strip())
            first_title = False
            i += 1
            continue
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=1)
            i += 1
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("- "):
            add_bullet(document, re.sub(r"\*\*(.*?)\*\*", r"\1", line[2:].strip()))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            add_bullet(document, re.sub(r"^\d+\.\s+", "", line), numbered=True)
            i += 1
            continue
        if line.startswith("|"):
            table_lines: list[str] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for item in table_lines:
                cells = [cell.strip() for cell in item.strip("|").split("|")]
                if not all(set(cell) <= {"-", ":", " "} for cell in cells):
                    rows.append(cells)
            if rows:
                table = document.add_table(rows=len(rows), cols=max(len(row) for row in rows))
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, value in enumerate(row):
                        cell = table.cell(r, c)
                        cell.text = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
                        if r == 0:
                            set_cell_shading(cell, "E0F5F0")
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            continue
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.add_run(re.sub(r"`([^`]+)`", r"\1", re.sub(r"\*\*(.*?)\*\*", r"\1", line)))
        i += 1

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("忆安智护 · 强网杯报名版 · 证据等级以运行门禁为准").font.size = Pt(8)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("markdown", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    render(args.markdown, args.output)
    print(f"wrote {args.output} ({args.output.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

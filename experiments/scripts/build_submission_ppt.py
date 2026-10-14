#!/usr/bin/env python3
"""Build the compact, evidence-bounded PPT used for competition submission.

The source deck is intentionally not modified. This output uses vector shapes
and text only so the official 20 MB limit is easy to verify.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


NAVY = RGBColor(15, 47, 61)
INK = RGBColor(27, 55, 69)
TEAL = RGBColor(18, 140, 140)
MINT = RGBColor(224, 245, 240)
GOLD = RGBColor(230, 181, 102)
CREAM = RGBColor(249, 247, 241)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(96, 117, 130)
LINE = RGBColor(207, 222, 225)
RED = RGBColor(180, 67, 54)


def set_bg(slide, color=CREAM):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def textbox(slide, x, y, w, h, text, size=18, color=INK, bold=False,
            align=PP_ALIGN.LEFT, font="Microsoft YaHei", margin=0.08):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = frame.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill, line=LINE, radius=True):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(0.8)
    return shape


def header(slide, section, page):
    textbox(slide, 0.65, 0.28, 10.7, 0.32, section.upper(), 9, TEAL, True)
    textbox(slide, 11.65, 0.28, 1.0, 0.32, f"{page:02d} / 12", 9, MUTED, False, PP_ALIGN.RIGHT)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(0.70), Inches(12.0), Inches(0.02)).fill.solid()


def title(slide, text, subtitle=None):
    textbox(slide, 0.65, 0.92, 11.8, 0.65, text, 28, NAVY, True)
    if subtitle:
        textbox(slide, 0.7, 1.58, 11.5, 0.42, subtitle, 12, MUTED)


def footer(slide, text="合成/合法脱敏数据 · 结果以服务端 trace 与证据门禁为准"):
    textbox(slide, 0.65, 7.08, 12.0, 0.24, text, 8.5, MUTED)


def card(slide, x, y, w, h, heading, body, accent=TEAL, body_size=13):
    rect(slide, x, y, w, h, WHITE)
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h)).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = accent
    slide.shapes[-1].line.fill.background()
    textbox(slide, x + 0.18, y + 0.15, w - 0.3, 0.30, heading, 14, NAVY, True)
    textbox(slide, x + 0.18, y + 0.53, w - 0.3, h - 0.62, body, body_size, INK)


def tag(slide, x, y, w, label, fill=MINT, color=TEAL):
    rect(slide, x, y, w, 0.32, fill, fill)
    textbox(slide, x, y + 0.01, w, 0.28, label, 9, color, True, PP_ALIGN.CENTER)


def pipeline(slide, labels, y=3.05):
    x = 0.75
    width = 2.22
    for index, label in enumerate(labels):
        rect(slide, x, y, width, 0.85, NAVY if index == 0 else WHITE, NAVY if index == 0 else LINE)
        textbox(slide, x + 0.06, y + 0.14, width - 0.12, 0.55, label, 14, WHITE if index == 0 else NAVY, True, PP_ALIGN.CENTER)
        if index < len(labels) - 1:
            textbox(slide, x + width, y + 0.18, 0.43, 0.42, "→", 24, GOLD, True, PP_ALIGN.CENTER)
        x += 2.62


def build(output: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank); set_bg(slide, NAVY)
    textbox(slide, 0.8, 0.75, 11.8, 0.35, "强网杯 · 信息安全创新作品", 12, GOLD, True)
    textbox(slide, 0.8, 1.55, 11.3, 0.9, "忆安智护", 42, WHITE, True)
    textbox(slide, 0.83, 2.52, 10.9, 0.62, "面向智慧养老的护理数据安全治理与可信辅助决策平台", 21, RGBColor(207, 232, 231), True)
    rect(slide, 0.85, 4.05, 11.45, 1.18, RGBColor(24, 73, 83), RGBColor(76, 138, 142))
    textbox(slide, 1.1, 4.32, 10.9, 0.62, "让护理数据记得清、查得到、防得住、说得明、删得净", 22, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, 0.85, 6.63, 8.0, 0.3, "现场演示主链：安全判定 → 因果证据 → 检索审计 → 删除验证", 11, RGBColor(207, 232, 231))
    textbox(slide, 10.0, 6.63, 2.3, 0.3, "报名版 v1 · 2026-10", 10, GOLD, True, PP_ALIGN.RIGHT)

    # 2. Problem
    slide = prs.slides.add_slide(blank); set_bg(slide); header(slide, "01 / 场景与问题", 2); title(slide, "智慧养老的难题，不是“有没有 AI”，而是“能不能被信任”")
    items = [
        ("碎片化", "护理记录来自文本、时间线、设备与图片，难以形成统一证据。"),
        ("高敏感", "健康、用药、行为与身份信息需要最小化展示和租户隔离。"),
        ("难追溯", "普通 RAG 只给答案，不能说明命中了哪一路、哪条证据。"),
        ("难撤回", "删除必须跨向量、图谱、时间线和缓存验证，而非只删一处。"),
    ]
    for i, (h, b) in enumerate(items):
        card(slide, 0.8 + (i % 2) * 6.0, 2.35 + (i // 2) * 1.72, 5.55, 1.35, h, b, [GOLD, TEAL, RED, NAVY][i])
    tag(slide, 0.82, 6.0, 2.05, "项目边界：辅助整理，不替代诊疗", RGBColor(255, 243, 225), RGBColor(150, 95, 12))
    footer(slide)

    # 3. Main chain
    slide = prs.slides.add_slide(blank); set_bg(slide); header(slide, "02 / 唯一比赛主链", 3); title(slide, "一条链路贯穿安全、证据、来源与删除")
    pipeline(slide, ["护理记录", "安全判定", "O-M-P-R\n证据", "人工复核", "检索审计", "删除验证"], 2.65)
    card(slide, 0.85, 4.35, 3.7, 1.45, "输入", "合成或合法脱敏护理记录\n每次请求绑定 trace_id 与作用域", TEAL)
    card(slide, 4.82, 4.35, 3.7, 1.45, "决策", "ASDF / CASIA / PCCM-S\n规则候选、模型裁决、证据校验", GOLD)
    card(slide, 8.79, 4.35, 3.7, 1.45, "证明", "RRF 来源审计 + 跨存储删除验证\n失败与降级如实保留", NAVY)
    footer(slide, "主链优先；多模态、三维记忆、冷热分层和研究 Agent 为受控扩展")

    # 4. Security
    slide = prs.slides.add_slide(blank); set_bg(slide); header(slide, "03 / 安全判定", 4); title(slide, "ASDF + CASIA + PCCM-S：先识别，再解释，再决定")
    card(slide, 0.8, 2.2, 3.75, 2.55, "ASDF · 对抗归一化", "空格、符号、同音、中文数字与 Base64 等候选绕过\n\n保留 normalization_steps、span_map、前后哈希和重检测结果", TEAL, 12.5)
    card(slide, 4.8, 2.2, 3.75, 2.55, "CASIA · 上下文判定", "使用版本化正负关键词和上下文窗口\n\n区分真实护理语境与攻击描述，记录命中词、权重和理由", GOLD, 12.5)
    card(slide, 8.8, 2.2, 3.75, 2.55, "PCCM-S · 证据校准", "融合规则、词典、上下文与安全模型证据\n\n输出 decision、置信度、早停/降级原因，不凭日志推断", NAVY, 12.5)
    tag(slide, 0.85, 5.38, 2.15, "结果等级", RGBColor(234, 244, 247), TEAL)
    textbox(slide, 3.2, 5.36, 8.6, 0.38, "verified  /  needs_review  /  blocked  /  insufficient_evidence", 16, NAVY, True)
    textbox(slide, 0.85, 6.12, 11.7, 0.42, "正式安全成绩只有在攻击、正常、歧义数据完成复核并通过 manifest 门禁后才写入报告。", 12, MUTED)
    footer(slide)

    # 5. Causality
    slide = prs.slides.add_slide(blank); set_bg(slide); header(slide, "04 / 因果证据", 5); title(slide, "O-M-P-R 不是一句生成答案，而是一组可回取的证据")
    fields = [("O", "Object", "患者/药物/疾病"), ("M", "Mediator", "机制或中介"), ("P", "Property", "状态变化"), ("R", "Result", "结果事件")]
    for i, (letter, name, body) in enumerate(fields):
        x = 0.8 + i * 3.05
        rect(slide, x, 2.25, 2.72, 2.0, WHITE)
        textbox(slide, x + 0.18, 2.45, 0.55, 0.55, letter, 29, TEAL, True, PP_ALIGN.CENTER)
        textbox(slide, x + 0.85, 2.46, 1.65, 0.28, name, 14, NAVY, True)
        textbox(slide, x + 0.85, 2.85, 1.65, 0.55, body, 12, INK)
        textbox(slide, x + 0.18, 3.55, 2.3, 0.4, "原文 span + rune 偏移", 10, MUTED)
    card(slide, 0.8, 4.75, 5.65, 1.25, "裁决链", "规则多候选 → 本地模型限时裁决 → 规则回退 → abstain", TEAL)
    card(slide, 6.85, 4.75, 5.65, 1.25, "拒答条件", "字段缺失、span 不在原文、否定/时序冲突或顺序错误时不补写", RED)
    footer(slide, "当前历史 Tuple F1 约 0.513 仅为问题基线；正式门槛通过前不宣称达标")

    # 6. Review and roles
    slide = prs.slides.add_slide(blank); set_bg(slide); header(slide, "05 / 人工复核", 6); title(slide, "把不确定性留给医生，把修改变成可追踪证据")
    card(slide, 0.8, 2.25, 3.75, 2.3, "待复核队列", "显示预测值、字段 span、否定/时序检查、模型层级和 trace。\n\n医生或评测角色才能提交合成数据纠错。", TEAL, 13)
    card(slide, 4.8, 2.25, 3.75, 2.3, "纠错记录", "保存原预测、修正值、证据 span、规则/模型版本。\n\n纠错不直接覆盖原结果，也不自动进入训练集。", GOLD, 13)
    card(slide, 8.8, 2.25, 3.75, 2.3, "四角色最小权限", "护工：记录与分析\n医生：证据复核\n家属：脱敏摘要\n老人：授权与删除请求", NAVY, 13)
    tag(slide, 0.82, 5.12, 2.35, "体验目标", RGBColor(235, 246, 240), TEAL)
    textbox(slide, 3.38, 5.11, 8.5, 0.42, "每个结论都能回答：来自哪里？为什么通过？谁改过？还能否撤回？", 17, NAVY, True)
    footer(slide)

    # 7. Retrieval
    slide = prs.slides.add_slide(blank); set_bg(slide); header(slide, "06 / 三存储与检索审计", 7); title(slide, "RRF 的价值不是“看起来更准”，而是每个结果都说得清")
    for i, (name, desc, color) in enumerate([
        ("Qdrant", "向量相似度\n排名 + 延迟", TEAL),
        ("Neo4j", "关系与实体\n排名 + 延迟", GOLD),
        ("TimescaleDB", "时间线与过滤\n排名 + 延迟", NAVY),
    ]):
        x = 0.9 + i * 4.1
        rect(slide, x, 2.22, 3.5, 1.55, WHITE)
        textbox(slide, x + 0.18, 2.45, 3.1, 0.35, name, 17, color, True)
        textbox(slide, x + 0.18, 2.92, 3.1, 0.45, desc, 12, INK)
    textbox(slide, 4.25, 4.05, 4.85, 0.42, "唯一 doc_id → 来源排名 → RRF 分数", 18, NAVY, True, PP_ALIGN.CENTER)
    card(slide, 0.9, 4.72, 5.55, 1.25, "成功契约", "retrieval_metadata 固定返回来源、排名、RRF 分数、单路状态、单路延迟、总延迟和 rrf_3_sources。", TEAL, 12)
    card(slide, 6.82, 4.72, 5.55, 1.25, "降级契约", "任一路为空、失败或超时，明确显示二路或向量回退；不把降级请求混入三路正式成绩。", RED, 12)
    footer(slide, "RRF MRR 提升仍需同数据、同 top-k、同清理流程和 readiness 证据后才能宣称")

    # 8. Deletion
    slide = prs.slides.add_slide(blank); set_bg(slide); header(slide, "07 / 删除验证", 8); title(slide, "用证明链验证“删干净”，而不是把它包装成机器遗忘")
    pipeline(slide, ["删除前\n计数/探针", "跨存储\n删除", "删除后\n计数/探针", "控制数据\n仍可检索", "幂等\n复删"], 2.45)
    card(slide, 0.82, 4.35, 3.7, 1.45, "验证对象", "Qdrant · Neo4j · TimescaleDB · session cache", TEAL)
    card(slide, 4.82, 4.35, 3.7, 1.45, "证明字段", "verification_id · trace_id · scope_hash · evidence_sha256", GOLD)
    card(slide, 8.82, 4.35, 3.7, 1.45, "统一口径", "隔离用户/会话跨存储删除验证\n不宣称参数级遗忘", NAVY)
    footer(slide, "删除失败、部分失败、补偿重试和控制作用域保留情况必须如实记录")

    # 9. Extensions
    slide = prs.slides.add_slide(blank); set_bg(slide); header(slide, "08 / 受控扩展", 9); title(slide, "能力可以扩展，但证据等级不能越级")
    ext = [
        ("多模态", "内存处理、哈希、大小限制；失败显示真实错误", "受控扩展"),
        ("记忆子图", "owner-scoped evidence API；进程内原型", "受控扩展"),
        ("冷热分层", "生命周期接口；迁移统计；非持久化原型", "受控扩展"),
        ("研究 Agent", "白名单来源、结构化引用、人工确认、离线回退", "受控扩展"),
    ]
    for i, (h, b, label) in enumerate(ext):
        x = 0.8 + (i % 2) * 6.0; y = 2.25 + (i // 2) * 1.55
        card(slide, x, y, 5.55, 1.18, h, b, MUTED, 12)
        tag(slide, x + 3.85, y + 0.16, 1.4, label, RGBColor(242, 242, 238), MUTED)
    textbox(slide, 0.85, 5.82, 11.6, 0.46, "它们服务于现场解释和后续研发，不参与安全、因果、RRF 正式指标。", 15, NAVY, True, PP_ALIGN.CENTER)
    footer(slide)

    # 10. Evaluation
    slide = prs.slides.add_slide(blank); set_bg(slide); header(slide, "09 / 实验与证据", 10); title(slide, "先固定数据和配置，再谈指标")
    headers = ["实验", "对照", "必须记录", "当前口径"]
    xs = [0.8, 2.55, 5.0, 9.0]
    widths = [1.55, 2.2, 3.75, 3.5]
    for x, w, h in zip(xs, widths, headers):
        rect(slide, x, 2.18, w, 0.45, NAVY, NAVY, False); textbox(slide, x, 2.25, w, 0.25, h, 11, WHITE, True, PP_ALIGN.CENTER)
    rows = [
        ("安全", "regex / ASDF / CASIA / PCCM-S", "分类型 P/R/F1、误伤、Wilson、延迟", "受控初步"),
        ("因果", "规则 / 3B / 裁决链", "Tuple、字段、否定、时序、p95", "历史基线"),
        ("RRF", "向量 / 过滤 / 三路", "MRR、Recall@5、来源与 readiness", "待正式运行"),
        ("删除", "目标 / 控制作用域", "计数、探针、幂等、证明哈希", "待 Docker 验证"),
    ]
    for r, row in enumerate(rows):
        y = 2.68 + r * 0.75
        for x, w, value in zip(xs, widths, row):
            rect(slide, x, y, w, 0.66, WHITE, LINE, False); textbox(slide, x + 0.04, y + 0.08, w - 0.08, 0.5, value, 10.5, INK if r < 2 else MUTED, r == 0, PP_ALIGN.CENTER)
    tag(slide, 0.84, 6.0, 2.15, "证据门禁", RGBColor(255, 243, 225), RGBColor(150, 95, 12))
    textbox(slide, 3.2, 5.99, 8.6, 0.42, "manifest · 配置指纹 · 模型版本 · 容器版本 · trace · 运行产物", 15, NAVY, True)
    footer(slide)

    # 11. Demo
    slide = prs.slides.add_slide(blank); set_bg(slide); header(slide, "10 / 现场演示", 11); title(slide, "90 秒只展示一条真实响应链")
    steps = [("01", "输入", "合成护理记录"), ("02", "安全", "ASDF/CASIA/PCCM-S"), ("03", "因果", "O-M-P-R span"), ("04", "检索", "三路来源审计"), ("05", "删除", "探针 + 证明哈希")]
    for i, (num, h, b) in enumerate(steps):
        x = 0.82 + i * 2.48
        rect(slide, x, 2.4, 2.15, 2.05, WHITE)
        textbox(slide, x + 0.12, 2.62, 0.62, 0.45, num, 22, TEAL, True)
        textbox(slide, x + 0.82, 2.65, 1.05, 0.3, h, 14, NAVY, True)
        textbox(slide, x + 0.18, 3.3, 1.75, 0.62, b, 13, INK, False, PP_ALIGN.CENTER)
    textbox(slide, 1.0, 5.18, 11.2, 0.55, "画面只出现真实服务响应；服务不可用时保留真实错误和降级状态。", 18, RED, True, PP_ALIGN.CENTER)
    textbox(slide, 1.1, 5.92, 11.0, 0.36, "录制前：Docker health check → fixture smoke → live smoke → release_guard", 12, MUTED, False, PP_ALIGN.CENTER)
    footer(slide)

    # 12. Close
    slide = prs.slides.add_slide(blank); set_bg(slide, NAVY)
    textbox(slide, 0.8, 0.75, 11.8, 0.35, "证据边界 · 可信交付", 12, GOLD, True)
    textbox(slide, 0.8, 1.55, 11.5, 0.75, "我们不承诺 AI 无所不能", 31, WHITE, True)
    textbox(slide, 0.83, 2.42, 11.4, 0.52, "我们承诺：每个决策有证据，每次降级可解释，每次删除可验证。", 20, RGBColor(207, 232, 231), True)
    rect(slide, 0.85, 3.65, 11.45, 1.45, RGBColor(24, 73, 83), RGBColor(76, 138, 142))
    textbox(slide, 1.12, 3.94, 10.9, 0.8, "已验证结果 / 受控初步结果 / 历史问题基线\n分级呈现，不用计划冒充成绩", 18, WHITE, True, PP_ALIGN.CENTER)
    textbox(slide, 0.85, 6.2, 10.6, 0.32, "源码：context-keeper-main · 主链：competition_console.html", 11, RGBColor(207, 232, 231))
    textbox(slide, 10.0, 6.2, 2.3, 0.32, "报名版 v1", 11, GOLD, True, PP_ALIGN.RIGHT)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.output)
    print(f"wrote {args.output} ({args.output.stat().st_size} bytes, 12 slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
